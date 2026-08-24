# -*- coding: utf-8 -*-
"""
Escribe la planilla madre ya ordenada, para usarla mientras la app se carga.

  python scripts/importar_plan.py   -> scripts/plan_semilla.json
  python scripts/plan_a_excel.py    -> Plan_Maestro_Limpio.xlsx

Cuatro hojas:
  Plan       una fila por actividad, con fecha, semana, dia, turno y HH.
             Es la misma informacion de los 56 bloques, pero en una tabla que se
             puede filtrar y sumar. Con filtros puestos y la fila de titulo fija.
  Carga      HH por dia contra las 344 disponibles, con los dias pasados marcados.
  Catalogo   las actividades y cuantas veces aparecen.
  Revisar    lo que quedo raro en la planilla original y hay que confirmar.
"""
import json, os, datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
SALIDA = os.path.join(RAIZ, 'Plan_Maestro_Limpio.xlsx')

datos = json.load(open(os.path.join(AQUI, 'plan_semilla.json'), encoding='utf-8'))
filas = datos['actividades']
TOPE = datos['hh_por_dia']

TITULO = Font(bold=True, color='FFFFFF')
FONDO = PatternFill('solid', fgColor='0D9488')
ROJO = PatternFill('solid', fgColor='FEE2E2')
GRIS = Font(color='64748B')


def encabezar(ws, cabeceras, anchos):
    ws.append(cabeceras)
    for i, (c, a) in enumerate(zip(cabeceras, anchos), start=1):
        celda = ws.cell(1, i)
        celda.font = TITULO
        celda.fill = FONDO
        celda.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = 'A2'


wb = Workbook()

# ---------------------------------------------------------------- hoja Plan
ws = wb.active
ws.title = 'Plan'
encabezar(ws, ['Fecha', 'Año', 'Semana', 'Día', 'Turno', 'Actividad', 'HH', 'Estado', 'Zona', 'OT', 'Nota'],
          [12, 7, 9, 12, 8, 42, 7, 13, 16, 12, 30])
for f in sorted(filas, key=lambda x: (x['fecha'], x['turno'] != 'Dia', x['orden'])):
    ws.append([
        datetime.date.fromisoformat(f['fecha']), f['anio'], f['semana'], f['dia'],
        'Día' if f['turno'] == 'Dia' else 'Noche', f['actividad'], f['hh'],
        'planificada', '', '', '',
    ])
for fila in ws.iter_rows(min_row=2, min_col=1, max_col=1):
    fila[0].number_format = 'DD-MM-YYYY'
ws.auto_filter.ref = ws.dimensions

# --------------------------------------------------------------- hoja Carga
carga = defaultdict(lambda: {'Dia': 0.0, 'Noche': 0.0})
for f in filas:
    carga[f['fecha']][f['turno']] += (f['hh'] or 0)

ws = wb.create_sheet('Carga')
encabezar(ws, ['Fecha', 'Semana', 'Día', 'HH día', 'HH noche', 'HH total',
               f'Disponible ({TOPE})', 'Libres'], [12, 9, 12, 10, 11, 10, 17, 10])
NOMBRE_DIA = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
for fecha in sorted(carga):
    d = datetime.date.fromisoformat(fecha)
    dia, noche = carga[fecha]['Dia'], carga[fecha]['Noche']
    total = dia + noche
    ws.append([d, d.isocalendar()[1], NOMBRE_DIA[d.weekday()], dia, noche, total, TOPE, TOPE - total])
    ws.cell(ws.max_row, 1).number_format = 'DD-MM-YYYY'
    if total > TOPE:
        for c in range(1, 9):
            ws.cell(ws.max_row, c).fill = ROJO
ws.auto_filter.ref = ws.dimensions

# ------------------------------------------------------------ hoja Catalogo
ws = wb.create_sheet('Catalogo')
encabezar(ws, ['Actividad', 'Veces en el plan'], [46, 16])
for c in datos['catalogo']:
    ws.append([c['actividad'], c['veces']])
ws.auto_filter.ref = ws.dimensions

# ------------------------------------------------------------- hoja Revisar
ws = wb.create_sheet('Revisar')
ws.column_dimensions['A'].width = 118
ws.append(['Lo que quedó raro en la planilla original y conviene confirmar'])
ws.cell(1, 1).font = TITULO
ws.cell(1, 1).fill = FONDO
diag = open(os.path.join(AQUI, 'plan_diagnostico.txt'), encoding='utf-8').read()
dentro = False
for linea in diag.splitlines():
    if linea.startswith('--- fechas'):
        dentro = True
        ws.append([''])
        ws.append([linea.strip('- ')])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        continue
    if linea.startswith('--- catalogo'):
        dentro = False
    if dentro and linea.strip():
        ws.append([linea.strip()])
        ws.cell(ws.max_row, 1).font = GRIS

ws.append([''])
ws.append([f'Días que pasan las {TOPE} HH disponibles: ver la hoja Carga, filas en rojo.'])
ws.cell(ws.max_row, 1).font = Font(bold=True)

wb.save(SALIDA)
print('escrito %s' % SALIDA)
print('%d filas en Plan, %d días en Carga, %d actividades en Catalogo'
      % (len(filas), len(carga), len(datos['catalogo'])))
